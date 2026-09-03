"""Excel -> list[list[str]] so the CSV path can take over."""
from datetime import date, datetime


def _cell_text(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def load_excel_rows(path, kind):
    if kind == "xls":
        return _load_xls(path)
    return _load_xlsx(path)


def _load_xlsx(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        best = []
        for ws in wb.worksheets:
            rows = []
            for r in ws.iter_rows(values_only=True):
                cells = [_cell_text(v) for v in r]
                if any(cells):
                    rows.append(cells)
            if len(rows) > len(best) and max((len(r) for r in rows), default=0) >= 2:
                best = rows
        return best
    finally:
        wb.close()


def _load_xls(path):
    import xlrd

    book = xlrd.open_workbook(path)
    best = []
    for sheet in book.sheets():
        rows = []
        for i in range(sheet.nrows):
            cells = []
            for j in range(sheet.ncols):
                c = sheet.cell(i, j)
                if c.ctype == xlrd.XL_CELL_DATE:
                    try:
                        cells.append(_cell_text(xlrd.xldate_as_datetime(c.value, book.datemode)))
                        continue
                    except (ValueError, OverflowError):
                        pass
                cells.append(_cell_text(c.value))
            if any(cells):
                rows.append(cells)
        if len(rows) > len(best) and sheet.ncols >= 2:
            best = rows
    return best
